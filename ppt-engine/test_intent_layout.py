import unittest
from copy import deepcopy
from intent_layout import ARCHETYPES,compose
from layout import validate,measure

ASSETS={'idol_stage':{},'idol_close':{}}
BASE={'title':'Festival economy','subtitle':'Community and commerce','body':'CAIT trade projection, not an audited total.',
      'quote':'A celebration supporting local businesses','items':[],'image_asset':None,
      'image_side':'right','background':'cream','ink':'ink','accent':'red','card':'gold','source':None}
PALETTE=[{'name':'cream','hex':'FFF8E1'},{'name':'ink','hex':'292929'},
         {'name':'red','hex':'A30000'},{'name':'gold','hex':'FFF4C9'}]

def slide(archetype,**changes):return dict(BASE,archetype=archetype,**changes)

class Archetypes(unittest.TestCase):
    def test_every_archetype_is_renderable_and_geometrically_clean(self):
        slides=[
          slide('title_hero',image_asset='idol_stage'),
          slide('big_stat',image_asset='idol_close',items=[{'heading':'Across India','value':'₹30,000 crore','text':'Festival business'}]),
          slide('card_grid',items=[{'heading':'Setup','value':'₹12,000 crore','text':'Decor and puja'},
                                   {'heading':'Food','value':'₹2,400 crore','text':'Sweets and vendors'}]),
          slide('split_image_text',image_asset='idol_stage'),
          slide('quote'),
          slide('timeline',items=[{'heading':'Prepare','value':None,'text':'Build pandals'},
                                   {'heading':'Celebrate','value':None,'text':'Attend local events'}]),
          slide('closing',items=[{'heading':'Times of India','value':None,
             'text':'https://timesofindia.indiatimes.com/city/pune/fuelled-by-faith-and-grandeur-ganeshotsav-a-rs-30k-cr-business-this-year/articleshow/123664315.cms'},
            {'heading':'CAIT','value':None,
             'text':'https://cait.in/ganesh-chaturthi-kickstarts-the-festive-season-sales-with-estimated-business-over-25000-crore-cait/'}])]
        plan,changes=compose({'palette':PALETTE,'slides':slides},ASSETS)
        self.assertEqual(len(plan['slides']),7)
        self.assertFalse(validate(plan))
        self.assertEqual([s['intent'] for s in plan['slides']],list(ARCHETYPES))
        self.assertTrue(all(e['type'] in {'text','box','image'} for s in plan['slides'] for e in s['elements']))
    def test_all_archetypes_export_native_and_pdf(self):
        import engine,json,tempfile
        from pathlib import Path
        from pptx import Presentation
        root=Path(__file__).parent
        assets={a['id']:a for a in json.loads((root/'assets.json').read_text())}
        slides=[slide('title_hero',image_asset='idol_stage'),
                slide('big_stat',image_asset='idol_close',items=[{'heading':'India','value':'₹30,000 crore','text':'Trade projection'}]),
                slide('card_grid',items=[{'heading':'Setup','value':'₹12,000 crore','text':'Decor and puja'},
                                         {'heading':'Food','value':'₹2,400 crore','text':'Vendors'}]),
                slide('split_image_text',image_asset='idol_close'),
                slide('quote'),
                slide('timeline',items=[{'heading':'Prepare','value':None,'text':'Build pandals'},
                                         {'heading':'Celebrate','value':None,'text':'Attend events'}]),
                slide('closing',items=[{'heading':'CAIT','value':None,'text':'https://cait.in/report'}])]
        plan,_=compose({'palette':PALETTE,'slides':slides},assets)
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);engine.render(plan,assets,p/'deck.pptx')
            images=engine.preview(p/'deck.pptx',p/'deck.pdf',p/'preview')
            self.assertEqual(len(images),7)
            self.assertEqual(len(Presentation(p/'deck.pptx').slides),7)
            self.assertGreater(sum(sh.has_text_frame for sl in Presentation(p/'deck.pptx').slides for sh in sl.shapes),20)
            self.assertEqual(engine.audit_rendered_text(plan,p/'deck.pdf'),[])
    def test_big_stat_subtitle_gets_anchored_panel(self):
        s=slide('big_stat',subtitle='CAIT 2025 trade projection',image_asset=None,
                items=[{'heading':None,'value':'₹30,000 crore','text':'Trade estimate'}])
        p,_=compose({'palette':PALETTE,'slides':[s]},ASSETS)
        elements=p['slides'][0]['elements']
        subtitle=next(e for e in elements if e.get('type')=='text' and 'CAIT 2025' in e['text'])
        self.assertTrue(any(e['type']=='box' and e['x']<=subtitle['x'] and e['x']+e['w']>=subtitle['x']+subtitle['w'] for e in elements))
        self.assertFalse(validate(p))
    def test_tall_cards_balance_content_vertically(self):
        s=slide('card_grid',items=[{'heading':'Setup','value':'₹12,000 crore','text':'Decor and puja'},
                                   {'heading':'Food','value':'₹2,400 crore','text':'Vendors'}])
        p,_=compose({'palette':PALETTE,'slides':[s]},ASSETS)
        elements=p['slides'][0]['elements']
        first_card=next(e for e in elements if e['type']=='box')
        children=[e for e in elements if e['type']=='text' and e['x']>=first_card['x'] and e['x']+e['w']<=first_card['x']+first_card['w'] and e['y']>=first_card['y']]
        self.assertGreaterEqual(max(e['y'] for e in children),first_card['y']+2.9)
        self.assertFalse(validate(p))
    def test_unfit_content_fails_without_lossy_rewrite(self):
        long='A long sentence about the festival and all its many unrelated details. '*30
        with self.assertRaisesRegex(ValueError,'text cannot fit'):
            compose({'palette':PALETTE,'slides':[slide('title_hero',title=long)]},ASSETS)
    def test_bad_model_geometry_is_not_part_of_intent(self):
        s=slide('big_stat',items=[{'heading':'India','value':'₹30,000 crore','text':''}])
        s['x']=-100;s['elements']=[{'type':'box','x':-100}]
        plan,_=compose({'palette':PALETTE,'slides':[s]},ASSETS)
        self.assertFalse(validate(plan))
        self.assertNotEqual(plan['slides'][0]['elements'][0].get('x'),-100)
    def test_unknown_asset_and_wrong_card_count_fail(self):
        with self.assertRaisesRegex(ValueError,'unknown image asset'):
            compose({'palette':PALETTE,'slides':[slide('split_image_text',image_asset='unknown')]},ASSETS)
        with self.assertRaisesRegex(ValueError,'2-4 cards'):
            compose({'palette':PALETTE,'slides':[slide('card_grid',items=[{'heading':'Only one'}])]},ASSETS)

if __name__=='__main__':unittest.main()
