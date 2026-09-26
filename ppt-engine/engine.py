#!/usr/bin/env python3
"""Model-directed presentation engine. Layout/creative choices are produced at runtime.
A small PPTX renderer accepts geometry but never picks a topic or layout itself.
"""
import argparse,base64,io,json,os,re,subprocess,sys,tempfile,time,urllib.request
from pathlib import Path
from PIL import Image,ImageOps
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

API='https://openrouter.ai/api/v1/chat/completions'
MAX_OUTPUT_TOKENS=6500

def chat(messages,model,base_url=API,max_tokens=MAX_OUTPUT_TOKENS):
    key=os.environ.get('OPENROUTER_API_KEY','')
    if not key:raise RuntimeError('OPENROUTER_API_KEY not set; provide it in the VPS process environment, not on the command line')
    body=json.dumps({'model':model,'messages':messages,'temperature':.5,'max_tokens':max_tokens},ensure_ascii=False).encode()
    req=urllib.request.Request(base_url,data=body,headers={'Authorization':'Bearer '+key,'Content-Type':'application/json','HTTP-Referer':'https://www.instinct.com','X-Title':'Presentation creative director'})
    last=None
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req,timeout=100) as res: ans=json.load(res)
            return ans['choices'][0]['message']['content'],ans.get('usage',{})
        except Exception as exc:
            last=exc
            if attempt<2:time.sleep(1.5*(attempt+1))
    raise RuntimeError(f'Model call failed after retries: {last}')

def parse_json(s):
    if not isinstance(s,str):raise ValueError('model response was not text')
    s=re.sub(r'^\s*```(?:json)?\s*|\s*```\s*$','',s.strip())
    return json.loads(s)

def prompt(brief,assets,reference):
    return '''You are a presentation creative director. Think through purpose, audience, story, visual metaphor, source reliability, and the reference design. You decide everything creative at runtime. The renderer only implements your design geometry. Respond with ONE JSON object, no markdown.

Canvas is 13.333 by 7.5 inches. You MUST include keys: rationale (string), palette (object of color names -> six-digit hex strings), slides (array, 4-12). Each slide has background (palette name or hex), intent (string), elements (array). Each element is ONE of:
- {"type":"text","text":"...","x":number,"y":number,"w":number,"h":number,"pt":number,"color":"palette key or hex","bold":bool,"align":"left|center|right"}
- {"type":"box","x":number,"y":number,"w":number,"h":number,"fill":"palette key or hex","rounded":bool}
- {"type":"image","asset":"asset ID from manifest","x":number,"y":number,"w":number,"h":number,"focus":[0.5,0.5]}
Geometry must fit canvas. Layer order is array order. Keep text editable, give whitespace and hierarchy. No inherited placeholders or stale navigation. Build a topic-appropriate deck, not a rigid template clone. Use real assets ONLY from the manifest; do not invent image URLs or imply that an illustrative image documents a number. Use supplied evidence only: do not invent figures, quotations, or source URLs. Put material source labels and uncertainties on slides. A sources page should carry URLs. Avoid minuscule type (body >=17pt, sources >=10pt). If brief asks another output format, stop and return {"unsupported_format":"..."}; this engine only produces PPTX and PDF.

USER BRIEF:\n'''+brief+'\n\nEVIDENCE / SOURCES:\n'+reference+'\n\nIMAGE MANIFEST:\n'+json.dumps(assets,ensure_ascii=False)

def color(v,palette):
    raw=palette.get(v,v).lstrip('#').upper()
    if not re.fullmatch(r'[A-F0-9]{6}',raw):raise ValueError('invalid color '+str(v))
    return RGBColor.from_string(raw)

def bounds(e):
    vals=[float(e[k]) for k in ('x','y','w','h')]
    x,y,w,h=vals
    if w<=0 or h<=0 or x<0 or y<0 or x+w>13.334 or y+h>7.501:raise ValueError('out-of-slide element '+str(vals))
    return [Inches(v) for v in vals]

def render(plan,assets,out):
    if 'unsupported_format' in plan:raise ValueError('Unsupported target format: '+str(plan['unsupported_format']))
    palette=plan['palette']; slides=plan['slides']
    if not 1<=len(slides)<=16:raise ValueError('invalid slide count')
    pres=Presentation();pres.slide_width=Inches(13.333);pres.slide_height=Inches(7.5)
    for slide in slides:
        sl=pres.slides.add_slide(pres.slide_layouts[6]);bg=sl.background.fill;bg.solid();bg.fore_color.rgb=color(slide['background'],palette)
        for e in slide['elements']:
            typ=e['type'];x,y,w,h=bounds(e)
            if typ=='box':
                sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e.get('rounded') else MSO_SHAPE.RECTANGLE,x,y,w,h);sh.fill.solid();sh.fill.fore_color.rgb=color(e['fill'],palette);sh.line.fill.background()
                if e.get('rounded'):sh.adjustments[0]=.09
            elif typ=='text':
                pt=float(e['pt']);
                if not 9<=pt<=120:raise ValueError('invalid font size')
                sh=sl.shapes.add_textbox(x,y,w,h);tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_top=tf.margin_bottom=0;tf.margin_left=tf.margin_right=Inches(.02)
                for i,line in enumerate(e['text'].split('\n')):
                    p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(0);p.line_spacing=1.08;p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[e.get('align','left')]
                    r=p.add_run();r.text=line;r.font.name='Inter';r.font.size=Pt(pt);r.font.bold=bool(e.get('bold'));r.font.color.rgb=color(e.get('color','222222'),palette)
            elif typ=='image':
                a=assets.get(e['asset']);
                if not a:raise ValueError('unknown image asset '+str(e['asset']))
                im=Image.open(a['file']).convert('RGB');center=e.get('focus',[.5,.5]);
                if len(center)!=2 or any(not 0<=float(v)<=1 for v in center):raise ValueError('invalid image focus')
                with io.BytesIO() as b:
                    ImageOps.fit(im,(max(100,round(w/914400*145)),max(100,round(h/914400*145))),centering=tuple(center)).save(b,format='JPEG',quality=88);b.seek(0);sl.shapes.add_picture(b,x,y,w,h)
            else:raise ValueError('unknown element type '+str(typ))
    pres.save(out)
    return len(slides)

def preview(pptx,pdf,folder):
    folder.mkdir(exist_ok=True,parents=True)
    cmd=['libreoffice','-env:UserInstallation=file:///tmp/lo-agent-deck-'+str(os.getpid()),'--headless','--convert-to','pdf','--outdir',str(folder),str(pptx)]
    subprocess.run(cmd,check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=90)
    generated=folder/(pptx.stem+'.pdf');generated.replace(pdf)
    import fitz
    d=fitz.open(pdf)
    images=[]
    for i,p in enumerate(d):
        pix=p.get_pixmap(matrix=fitz.Matrix(.8,.8),alpha=False)
        f=folder/f'slide-{i+1:02}.png';pix.save(f);images.append(f)
    return images

def critic_prompt(plan,images):
    content=[{'type':'text','text':'Critique this rendered deck against its brief. Look at EVERY image for clipping, overlap, blank/sloppy composition, readability, mismatched imagery, weak story, unsupported claims, and misleading citations. If flaws exist return JSON {"approved":false,"issues":[{"slide":1,"problem":"..."}],"revised_plan":<FULL revised plan schema>}. If excellent return {"approved":true,"issues":[]}. Make substantive redesign choices yourself; do not merely replace placeholders. Current plan: '+json.dumps(plan,ensure_ascii=False)}]
    for f in images:
        b64=base64.b64encode(f.read_bytes()).decode();content.append({'type':'image_url','image_url':{'url':'data:image/png;base64,'+b64,'detail':'low'}})
    return content

def run(args):
    brief=Path(args.brief).read_text();reference=Path(args.sources).read_text() if args.sources else ''
    manifest=json.loads(Path(args.assets).read_text());assets={a['id']:a for a in manifest}
    for a in assets.values():
        if not Path(a['file']).is_file():raise FileNotFoundError(a['file'])
    out=Path(args.output);out.mkdir(parents=True,exist_ok=True)
    response,usage=chat([{'role':'system','content':'You are the creative lead, responsible for design decisions, not a slot filler.'},{'role':'user','content':prompt(brief,manifest,reference)}],args.model,args.api_url)
    plan=parse_json(response);(out/'plan-0.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2))
    log=[{'stage':'planner','usage':usage}]
    for iteration in range(args.max_revisions+1):
        pptx=out/'deck.pptx';pdf=out/'deck.pdf';render(plan,assets,pptx)
        images=preview(pptx,pdf,out/f'preview-{iteration}')
        if not args.vision_model:
            (out/'run-log.json').write_text(json.dumps({'iterations':log,'visual_qa':'UNVERIFIED: no vision model supplied; human review required'},indent=2))
            print('Generated but visually unverified; supply --vision-model to run critic loop.');return 2
        raw,usage=chat([{'role':'system','content':'You are a rigorous presentation critic. Return strict JSON.'},{'role':'user','content':critic_prompt(plan,images)}],args.vision_model,args.api_url,max_tokens=7000)
        result=parse_json(raw);log.append({'stage':'critic','iteration':iteration,'usage':usage,'issues':result.get('issues',[])})
        (out/f'critique-{iteration}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
        if result.get('approved') is True:
            (out/'run-log.json').write_text(json.dumps({'iterations':log,'visual_qa':'model-inspected all slide previews; review PDF as final viewer check'},indent=2));print('Approved',pptx,pdf);return 0
        if iteration>=args.max_revisions:break
        if 'revised_plan' not in result:raise ValueError('critic rejected without full revised_plan')
        plan=result['revised_plan'];(out/f'plan-{iteration+1}.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2))
    (out/'run-log.json').write_text(json.dumps({'iterations':log,'visual_qa':'FAILED: unresolved critic issues'},indent=2))
    print('Critic did not approve; files are drafts, not deliverable',file=sys.stderr);return 3

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--brief',required=True);ap.add_argument('--assets',required=True);ap.add_argument('--sources');ap.add_argument('--output',required=True);ap.add_argument('--model',default='google/gemini-2.5-flash');ap.add_argument('--vision-model',default='google/gemini-2.5-flash');ap.add_argument('--api-url',default=API,help='OpenAI-compatible test endpoint override; keep default on VPS');ap.add_argument('--max-revisions',type=int,default=2)
    try:sys.exit(run(ap.parse_args()))
    except Exception as exc: print('ERROR:',exc,file=sys.stderr);sys.exit(1)
